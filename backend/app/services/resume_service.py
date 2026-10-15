import json
import logging
from io import BytesIO
from pathlib import Path
from typing import Any

logger = logging.getLogger(__name__)

import httpx
from docx import Document
from fastapi import HTTPException, UploadFile
from pypdf import PdfReader

from app.core.config import settings

ALLOWED_EXTENSIONS = {".pdf", ".docx", ".txt"}


def extract_resume_text(filename: str, content: bytes) -> str:
    extension = Path(filename).suffix.lower()
    if extension not in ALLOWED_EXTENSIONS:
        raise HTTPException(status_code=415, detail="Only PDF, DOCX, and TXT resumes are supported.")
    try:
        if extension == ".pdf":
            try:
                reader = PdfReader(BytesIO(content))
                text = "\n".join(page.extract_text() or "" for page in reader.pages)
            except Exception:
                # Fallback if raw text or demo stream was passed
                text = content.decode("utf-8", errors="ignore")
        elif extension == ".docx":
            document = Document(BytesIO(content))
            text = "\n".join(paragraph.text for paragraph in document.paragraphs)
            text += "\n" + "\n".join(cell.text for table in document.tables for row in table.rows for cell in row.cells)
        else:
            text = content.decode("utf-8", errors="ignore")
    except Exception as exc:
        raise HTTPException(status_code=422, detail="The uploaded resume could not be read.") from exc
    cleaned = "\n".join(line.strip() for line in text.splitlines() if line.strip())
    if len(cleaned) < 80:
        raise HTTPException(status_code=422, detail="The resume does not contain enough readable text.")
    return cleaned


def build_prompt(job_description: str, resume: str, sections: str = "summary,experience,skills", tone: str = "impactful") -> str:
    return f"""You are an elite executive resume writer, hiring manager, and ATS (Applicant Tracking System) optimization specialist.

Task:
Tailor the candidate resume to match the job description accurately and compellingly, without inventing facts.
Preserve every truthful achievement, employer, title, date, education item, and metric.
Focus customization on the requested sections: {sections}.
Adopt a professional tone: {tone}.
Improve relevance by reordering content, emphasizing matching skills, and rewriting bullets using the Google XYZ formula ("Accomplished [X], as measured by [Y], by doing [Z]") and strong action verbs.
Never fabricate skills, certifications, employers, metrics, or experiences that are not present in the source resume.

Return ONLY valid JSON matching this exact structure:
{{
  "summary": "A concise 2-sentence explanation of the tailoring strategy applied",
  "match_score": 85,
  "matched_keywords": ["keyword 1", "keyword 2", "keyword 3"],
  "missing_keywords": ["missing skill 1", "missing qualification 2"],
  "ats_breakdown": {{
    "overall": 88,
    "keyword_match": 84,
    "action_verbs": 92,
    "formatting": 95,
    "recommendations": [
      "Explicitly include cloud architecture metrics in the summary",
      "Add certifications if already achieved to beat hard keyword filters",
      "Keep bullet points between 1-2 lines for optimal ATS readability"
    ]
  }},
  "cover_letter": "A polished, tailored 3-paragraph cover letter for this role matching the candidate's verified achievements.",
  "resume_markdown": "The complete tailored resume in clean, professional Markdown with consistent # and ## headers and bullet points."
}}

JOB DESCRIPTION:
{job_description}

SOURCE RESUME:
{resume}
"""


def generate_docx_stream(markdown_text: str) -> BytesIO:
    doc = Document()
    for raw_line in markdown_text.splitlines():
        line = raw_line.strip()
        if not line:
            continue
        if line.startswith("# "):
            doc.add_heading(line[2:].strip(), level=1)
        elif line.startswith("## "):
            doc.add_heading(line[3:].strip(), level=2)
        elif line.startswith("### "):
            doc.add_heading(line[4:].strip(), level=3)
        elif line.startswith("- ") or line.startswith("* "):
            doc.add_paragraph(line[2:].strip(), style="List Bullet")
        else:
            # Clean bold markers for cleaner word doc look
            clean_text = line.replace("**", "")
            doc.add_paragraph(clean_text)
    buffer = BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer


async def customize_resume(
    job_description: str,
    resume_file: UploadFile,
    sections: str = "summary,experience,skills",
    tone: str = "impactful",
) -> dict[str, Any]:
    if not settings.gemini_api_key:
        raise HTTPException(status_code=503, detail="Gemini is not configured. Add GEMINI_API_KEY to backend/.env.")
    if not job_description.strip():
        raise HTTPException(status_code=400, detail="Job description is required.")
    
    content = await resume_file.read()
    if len(content) > settings.max_upload_size_mb * 1024 * 1024:
        raise HTTPException(status_code=413, detail="Resume file is too large.")
    
    filename = resume_file.filename or "resume.pdf"
    file_size_bytes = len(content)
    resume_text = extract_resume_text(filename, content)
    
    prompt = build_prompt(job_description, resume_text, sections, tone)
    payload = {
        "contents": [{"parts": [{"text": prompt}]}],
        "generationConfig": {"temperature": 0.25, "responseMimeType": "application/json"},
    }
    models_to_try = [settings.gemini_model]
    for fallback in ["gemini-3-flash-preview", "gemini-3.5-flash", "gemini-3.6-flash", "gemini-3.1-flash-lite"]:
        if fallback not in models_to_try:
            models_to_try.append(fallback)

    data = None
    last_error_message = None

    async with httpx.AsyncClient(timeout=90) as client:
        for model_name in models_to_try:
            url = f"https://generativelanguage.googleapis.com/v1beta/models/{model_name}:generateContent"
            try:
                response = await client.post(url, params={"key": settings.gemini_api_key}, json=payload)
                if response.status_code == 200:
                    data = response.json()
                    break

                error_body = response.text
                logger.warning(
                    "Gemini API error %s on model '%s': %s",
                    response.status_code,
                    model_name,
                    error_body,
                )
                last_error_message = f"Model {model_name} returned {response.status_code}"

                if response.status_code in (401, 403):
                    raise HTTPException(
                        status_code=502,
                        detail="Invalid Gemini API key or permission denied.",
                    )
            except httpx.RequestError as exc:
                logger.warning("Network error calling Gemini model '%s': %s", model_name, exc)
                last_error_message = f"Network connection error: {exc}"

    if not data:
        raise HTTPException(
            status_code=502,
            detail=f"Gemini could not process this request. ({last_error_message or 'All model attempts failed'})",
        )

    try:
        raw = data["candidates"][0]["content"]["parts"][0]["text"]
        result = json.loads(raw)
        if not isinstance(result, dict) or not result.get("resume_markdown"):
            raise ValueError("Incomplete Gemini response")

        ats_raw = result.get("ats_breakdown") or {}
        ats_breakdown = {
            "overall": max(0, min(100, int(ats_raw.get("overall", result.get("match_score", 85))))),
            "keyword_match": max(0, min(100, int(ats_raw.get("keyword_match", 80)))),
            "action_verbs": max(0, min(100, int(ats_raw.get("action_verbs", 85)))),
            "formatting": max(0, min(100, int(ats_raw.get("formatting", 90)))),
            "recommendations": [str(r) for r in ats_raw.get("recommendations", [])][:5],
        }

        matched_kw = result.get("matched_keywords") or result.get("keywords") or []
        missing_kw = result.get("missing_keywords") or []

        return {
            "summary": str(result.get("summary", "Resume tailored strategically to the target role.")),
            "match_score": max(0, min(100, int(result.get("match_score", ats_breakdown["overall"])))),
            "keywords": [str(item) for item in matched_kw][:15],
            "matched_keywords": [str(item) for item in matched_kw][:15],
            "missing_keywords": [str(item) for item in missing_kw][:12],
            "ats_breakdown": ats_breakdown,
            "cover_letter": str(result.get("cover_letter", "")),
            "resume_markdown": str(result["resume_markdown"]),
            "original_text": resume_text,
            "file_info": {
                "name": filename,
                "size_bytes": file_size_bytes,
                "size_formatted": f"{(file_size_bytes / 1024):.1f} KB" if file_size_bytes < 1024 * 1024 else f"{(file_size_bytes / (1024 * 1024)):.2f} MB",
            },
        }
    except (KeyError, IndexError, TypeError, ValueError, json.JSONDecodeError) as exc:
        raise HTTPException(status_code=502, detail="Gemini returned an invalid customization response.") from exc
